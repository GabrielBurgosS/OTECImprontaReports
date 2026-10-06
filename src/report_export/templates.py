"""Catalog of report templates and automatic detection of what each column holds."""
import json
import logging
import re
import shutil
import unicodedata
import warnings
from dataclasses import asdict, dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, range_boundaries
from openpyxl.worksheet.formula import ArrayFormula

from src.errors import TemplateError
from src.paths import bundled_templates_dir, templates_dir

logger = logging.getLogger(__name__)

CATALOG_FILE = 'plantillas.json'
PLAIN_TEMPLATE_ID = 'datos'

# Column roles
ID, NAME, EMAIL, PARTICIPATION, GRADE, ATTENDANCE, FORMULA, MANUAL = (
    'id', 'name', 'email', 'participation', 'grade', 'attendance', 'formula', 'manual')
DATA_ROLES = (ID, NAME, EMAIL, PARTICIPATION, GRADE)

ROLE_LABELS = {
    ID: 'ID', NAME: 'nombre', EMAIL: 'correo', PARTICIPATION: '% de participación',
    GRADE: 'nota', ATTENDANCE: 'asistencia (a mano)', FORMULA: 'calculada por Excel',
    MANUAL: 'se llena a mano',
}

BUILTIN_TEMPLATES = [
    {'id': 'inicial', 'name': 'Reporte Inicial', 'file': 'Reporte Inicial.xlsx', 'prefix': 'Reporte Inicial'},
    {'id': 'semanal', 'name': 'Reporte Semanal', 'file': 'Reporte Semanal.xlsx', 'prefix': 'Reporte Semanal'},
    {'id': 'part', 'name': 'Reporte de Participación', 'file': 'Reporte Part.xlsx',
     'prefix': 'Reporte de Participación'},
    {'id': 'final', 'name': 'Reporte Final', 'file': 'Reporte Final.xlsx', 'prefix': 'Reporte Final'},
]
PLAIN_TEMPLATE = {'id': PLAIN_TEMPLATE_ID, 'name': 'Datos completos (sin formato)', 'file': '',
                  'prefix': 'Datos Completos'}


def normalize(text) -> str:
    text = unicodedata.normalize('NFKD', str(text or '')).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', text).strip().lower()


@dataclass
class TemplateColumn:
    header: str
    role: str
    group: str = ''


@dataclass
class TemplateLayout:
    sheet: str
    table: str
    columns: list[TemplateColumn]

    def headers(self, role: str) -> list[str]:
        return [c.header for c in self.columns if c.role == role]

    @property
    def has_attendance(self) -> bool:
        return any(c.role == ATTENDANCE for c in self.columns)

    def summary(self) -> str:
        filled = sum(c.role in DATA_ROLES for c in self.columns)
        manual = sum(c.role in (MANUAL, ATTENDANCE) for c in self.columns)
        formulas = sum(c.role == FORMULA for c in self.columns)
        parts = [f'llena {filled} columnas']
        if formulas:
            parts.append(f'{formulas} calculadas')
        if manual:
            parts.append(f'{manual} a mano')
        return ' · '.join(parts)


@dataclass
class TemplateInfo:
    id: str
    name: str
    file: str
    prefix: str
    builtin: bool = False

    @property
    def is_plain(self) -> bool:
        return self.id == PLAIN_TEMPLATE_ID

    @property
    def path(self) -> Path:
        return getattr(self, '_path', None) or templates_dir() / self.file


def _column_role(header: str, group: str, first_value) -> str:
    h, g = normalize(header), normalize(group)
    if isinstance(first_value, ArrayFormula) or (isinstance(first_value, str) and first_value.startswith('=')):
        return FORMULA
    if h in ('id', 'rut', 'run', 'numero de id', 'n de id'):
        return ID
    if h in ('nombre', 'nombre completo', 'nombres'):
        return NAME
    if 'correo' in h or 'email' in h or 'e-mail' in h:
        return EMAIL
    if h.startswith('%') and ('part' in h or 'avance' in h):
        return PARTICIPATION
    if 'asistencia' in g:
        return ATTENDANCE
    if 'nota' in g or h.startswith(('prueba', 'nota', 'evaluacion', 'test', 'control')):
        return GRADE
    return MANUAL


def analyze_template(path: Path) -> TemplateLayout:
    """Read the template and classify each column of its Excel table."""
    if not path.exists():
        raise TemplateError(f'Template file not found: {path}',
                            f'No se encontró el archivo de plantilla "{path.name}".')
    try:
        with warnings.catch_warnings():
            warnings.simplefilter('ignore')  # openpyxl warns about extensions it would drop on save
            workbook = openpyxl.load_workbook(path)
    except Exception as e:
        raise TemplateError(f'Cannot open template {path}: {e}',
                            f'No se pudo abrir la plantilla "{path.name}". ¿Es un archivo de Excel válido?') from e

    for sheet in workbook.worksheets:
        for table in sheet.tables.values():
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            columns = []
            group = ''
            for col in range(min_col, max_col + 1):
                header = sheet.cell(min_row, col).value
                if min_row > 1 and sheet.cell(min_row - 1, col).value:
                    group = str(sheet.cell(min_row - 1, col).value)
                first_value = sheet.cell(min_row + 1, col).value if max_row > min_row else None
                columns.append(TemplateColumn(str(header), _column_role(str(header), group, first_value), group))
            layout = TemplateLayout(sheet.title, table.displayName, columns)
            if not (layout.headers(NAME) or layout.headers(EMAIL)):
                continue
            return layout
    raise TemplateError(f'No usable table in {path}',
                        f'La plantilla "{path.name}" debe tener una tabla de Excel (Insertar → Tabla) '
                        'con al menos una columna "Nombre" o "Dirección de correo".')


class TemplateCatalog:
    """Templates available to the user, stored in the user templates folder."""

    def __init__(self, folder: Path | None = None):
        self.folder = folder or templates_dir()
        self.catalog_path = self.folder / CATALOG_FILE
        if not self.catalog_path.exists():
            self.restore_builtin()
        self.templates = self._load()

    def _load(self) -> list[TemplateInfo]:
        try:
            entries = json.loads(self.catalog_path.read_text(encoding='utf-8'))
        except (OSError, json.JSONDecodeError):
            logger.exception('Template catalog unreadable; restoring the built-in templates')
            self.restore_builtin()
            entries = json.loads(self.catalog_path.read_text(encoding='utf-8'))
        templates = [TemplateInfo(**e) for e in entries]
        for template in templates:
            template._path = self.folder / template.file
        return [TemplateInfo(**PLAIN_TEMPLATE, builtin=True)] + [t for t in templates if not t.is_plain]

    def _save(self) -> None:
        entries = [asdict(t) for t in self.templates if not t.is_plain]
        self.catalog_path.write_text(json.dumps(entries, ensure_ascii=False, indent=2), encoding='utf-8')

    def restore_builtin(self) -> None:
        """Copy the original templates again, keeping any template the user added."""
        custom = []
        if self.catalog_path.exists():
            try:
                custom = [e for e in json.loads(self.catalog_path.read_text(encoding='utf-8'))
                          if not e.get('builtin')]
            except (OSError, json.JSONDecodeError):
                pass
        for entry in BUILTIN_TEMPLATES:
            shutil.copy2(bundled_templates_dir() / entry['file'], self.folder / entry['file'])
        entries = [dict(e, builtin=True) for e in BUILTIN_TEMPLATES] + custom
        self.catalog_path.write_text(json.dumps(entries, ensure_ascii=False, indent=2), encoding='utf-8')
        self.templates = self._load() if hasattr(self, 'templates') else []
        logger.info('Plantillas originales restauradas')

    def get(self, template_id: str) -> TemplateInfo:
        for template in self.templates:
            if template.id == template_id:
                return template
        raise TemplateError(f'Unknown template id {template_id}',
                            'La plantilla elegida ya no existe. Elige otra en la lista.')

    def layout(self, template_id: str) -> TemplateLayout | None:
        template = self.get(template_id)
        return None if template.is_plain else analyze_template(template.path)

    def add(self, source: Path, name: str, prefix: str | None = None) -> TemplateInfo:
        """Copy a new template file into the catalog after validating it."""
        analyze_template(source)
        slug = re.sub(r'[^a-z0-9]+', '-', normalize(name)).strip('-') or 'plantilla'
        template_id, counter = slug, 2
        while any(t.id == template_id for t in self.templates):
            template_id, counter = f'{slug}-{counter}', counter + 1
        file_name = f'{template_id}.xlsx'
        shutil.copy2(source, self.folder / file_name)
        template = TemplateInfo(template_id, name.strip(), file_name, (prefix or name).strip())
        self.templates.append(template)
        self._save()
        logger.info('Plantilla agregada: %s (%s)', template.name, source.name)
        return template

    def remove(self, template_id: str) -> None:
        template = self.get(template_id)
        if template.builtin:
            raise TemplateError(f'Cannot remove builtin template {template_id}',
                                'Las plantillas originales no se pueden eliminar.')
        self.templates.remove(template)
        self._save()
        try:
            template.path.unlink()
        except OSError:
            pass


@dataclass
class GradeMapping:
    """Which platform quiz goes in each grade column of a template."""
    mapping: dict[str, str | None] = field(default_factory=dict)
    ambiguous: list[str] = field(default_factory=list)


def suggest_grade_mapping(grade_headers: list[str], quizzes: list[str],
                          remembered: dict | None = None) -> GradeMapping:
    """Match template grade columns with quizzes by name; flag the ones that need the user."""
    remembered = remembered or {}
    result = GradeMapping()
    for header in grade_headers:
        if header in remembered and (remembered[header] is None or remembered[header] in quizzes):
            result.mapping[header] = remembered[header]
            continue
        h = normalize(header)
        if 'diag' in h:
            candidates = [q for q in quizzes if 'diag' in normalize(q)]
        elif 'final' in h:
            candidates = [q for q in quizzes if 'final' in normalize(q)]
        else:
            candidates = [q for q in quizzes if normalize(q) == h] or \
                         [q for q in quizzes if h in normalize(q) or normalize(q) in h]
        if len(candidates) == 1:
            result.mapping[header] = candidates[0]
        else:
            result.mapping[header] = None
            result.ambiguous.append(header)
    return result
